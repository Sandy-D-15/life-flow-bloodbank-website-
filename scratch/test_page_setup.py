import os
import docx
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Section 0: Cover
s0 = doc.sections[0]
s0.different_first_page_header_footer = False
# No header or footer added to s0
doc.add_paragraph('Cover Page (Unnumbered)')

# Section 1: Front Matter
s1 = doc.add_section(WD_SECTION.NEW_PAGE)
s1.different_first_page_header_footer = False
s1.header.is_linked_to_previous = False
s1.footer.is_linked_to_previous = False

sectPr1 = s1._sectPr
pgNumType1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="2" w:fmt="lowerRoman"/>')
sectPr1.append(pgNumType1)

p_f1 = s1.footer.paragraphs[0]
p_f1.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
p_f1._p.append(fld1)

doc.add_paragraph('Certificate Page (should show ii)')
doc.add_page_break()
doc.add_paragraph('Abstract Page (should show iii)')

# Section 2: Body
s2 = doc.add_section(WD_SECTION.NEW_PAGE)
s2.different_first_page_header_footer = False
s2.header.is_linked_to_previous = False
s2.footer.is_linked_to_previous = False

sectPr2 = s2._sectPr
pgNumType2 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="decimal"/>')
sectPr2.append(pgNumType2)

p_h2 = s2.header.paragraphs[0]
p_h2.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
r = p_h2.add_run('LifeFlow Blood Bank Management System')
r.font.size = docx.shared.Pt(8.5)

p_f2 = s2.footer.paragraphs[0]
p_f2.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
fld2 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
p_f2._p.append(fld2)

doc.add_paragraph('Chapter 1 (should show 1)')
doc.add_page_break()
doc.add_paragraph('Chapter 1 Continued (should show 2)')

os.makedirs('scratch', exist_ok=True)
doc.save('scratch/test_pagination.docx')
print('Saved test_pagination.docx')
